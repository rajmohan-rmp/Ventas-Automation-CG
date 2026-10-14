#!/usr/bin/env python3
"""
ventas_quantity_generator.py — a second, log-scale QUANTITY dashboard (1 slide).

Charts, on a logarithmic Y axis, the daily quantity for:
  Primary Sales  -> TOTAL_QUANTITY
  POS Sales      -> TOTAL_QUANTITY
  Stock Transfer -> REQUESTED_QUANTITY
summed per date across order types, over the same 7-day window as the main deck.

Usage:
  python ventas_quantity_generator.py --csv-file a.csv b.csv ... --date YYYY-MM-DD --output qty.pptx [--png qty.png]
"""
import argparse, sys
from pathlib import Path
from datetime import datetime, timedelta

import ventas_slide_generator as G   # reuse CSV parsing / helpers (same folder)
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

BG="#06224A"; CHART_BG2="#0E3A6E"; HDR="#0E518C"

def _to_float(v):
    try: return float(str(v).strip().replace(",",""))
    except: return 0.0

def _compact(v):
    """Short label so wide numbers don't overlap on narrow log panels: 1.2K / 3.4M / 1.3B."""
    a=abs(v)
    if a>=1e9: return f"{v/1e9:.1f}B"
    if a>=1e6: return f"{v/1e6:.1f}M"
    if a>=1e3: return f"{v/1e3:.1f}K"
    return f"{v:,.0f}" if abs(v-round(v))<1e-9 else f"{v:,.2f}"

def _cfg():
    try:
        import yaml
        return yaml.safe_load((Path(__file__).parent/"config.yaml").read_text(encoding="utf-8")) or {}
    except Exception:
        return {}

def _window(report_date):
    # Mirror ventas_slide_generator.py EXACTLY so both decks in the same email
    # always cover the identical dates: last 7 days, excluding the report day
    # (report_date - 1 back through report_date - 7). The main generator hardcodes
    # this and does not read config, so neither do we.
    end = datetime.strptime(report_date,"%Y-%m-%d") - timedelta(days=1)
    return [(end-timedelta(days=6-i)).strftime("%Y-%m-%d") for i in range(7)]

# metric spec: classify label -> (display title, quantity column, bar colour)
SPEC = {
    "Primary Sales Summary": ("Primary Sales — Total Quantity", "TOTAL_QUANTITY", "#2ECC71"),
    "POS Sales Summary":     ("POS Sales — Total Quantity",     "TOTAL_QUANTITY", "#1172B8"),
    "Stock Transfer Summary":("Stock Transfer — Requested Qty", "REQUESTED_QUANTITY", "#F1C40F"),
}
ORDER = ["Primary Sales Summary","POS Sales Summary","Stock Transfer Summary"]

def _collect(csv_paths):
    """Return {classify_label: {date: quantity}} using the configured column."""
    out={k:{} for k in SPEC}
    for p in csv_paths:
        key=G._classify(p)
        if key not in SPEC: continue
        col=SPEC[key][1]
        rows=G._read_csv_rows(p)
        if not rows: continue
        have = col in rows[0]
        for r in rows:
            d=G._date_key(r)
            if not d: continue
            out[key][d]=out[key].get(d,0.0)+ (_to_float(r.get(col)) if have else 0.0)
    return out

def render_png(data, report_date, out_png, dpi=150):
    from matplotlib.patches import FancyBboxPatch
    CARD="#0E3A6E"; INK="white"; SUB="#9FB6D6"; GRID="#2C4E7A"; UP="#2ECC71"; DOWN="#E74C3C"
    dates=_window(report_date)
    short=lambda d: datetime.strptime(d,"%Y-%m-%d").strftime("%d-%b")
    customer=G.customer_name(_cfg())   # single source: branding.country / country_code
    x=np.arange(len(dates))
    fig=plt.figure(figsize=(13.333,7.5),dpi=dpi); fig.patch.set_facecolor(BG)
    bg=fig.add_axes([0,0,1,1]); bg.set_axis_off(); bg.set_xlim(0,1); bg.set_ylim(0,1)
    bg.add_patch(Rectangle((0,0.93),1,0.07,color=HDR,zorder=1))
    bg.text(0.015,0.967,"VENTAS QUANTITY DASHBOARD",color=INK,fontsize=20,fontweight="bold",va="center",zorder=3)
    bg.text(0.015,0.94,"%s    |    Report Date: %s    |    Last 7 Days    |    LOG SCALE"
            %(customer, G._short_date_long(report_date)),
            color=SUB,fontsize=9.5,va="center",zorder=3)
    bg.text(0.985,0.962,"6D Technologies",color=INK,fontsize=12,fontweight="bold",ha="right",va="center",zorder=3)

    M=0.02; colw=(1-4*M)/3
    TOP_Y0=0.725; TOP_Y1=0.885       # summary box
    BOT_Y0=0.05;  BOT_Y1=0.705       # graph box
    for i,key in enumerate(ORDER):
        title,col,color=SPEC[key]
        vals=[data.get(key,{}).get(d,0.0) for d in dates]
        x0=M+i*(colw+M); px=0.016
        pos=[v for v in vals if v>0]
        total=sum(vals); avg=total/len(vals) if vals else 0
        peak=max(vals) if vals else 0; peakday=short(dates[vals.index(peak)]) if peak>0 else "-"
        change=(vals[-1]-avg)/avg*100 if avg else 0           # latest day vs 7-day average
        arrow="▲" if change>=0 else "▼"; acol=UP if change>=0 else DOWN

        # ---- top summary box ----
        top=FancyBboxPatch((x0,TOP_Y0),colw,TOP_Y1-TOP_Y0,boxstyle="round,pad=0.002,rounding_size=0.012",
                           fc=CARD,ec=GRID,lw=1.2,zorder=2); bg.add_patch(top)
        bg.text(x0+colw/2,TOP_Y1-0.026,title,color=INK,fontsize=12,fontweight="bold",ha="center",va="center",zorder=4)
        lab_y=TOP_Y1-0.07; val_y=TOP_Y1-0.115; sub_y=TOP_Y1-0.148
        lx=x0+px; cxm=x0+colw/2; rx=x0+colw-px
        bg.text(lx,lab_y,"7-DAY TOTAL",color=SUB,fontsize=8.5,ha="left",va="center",zorder=4)
        bg.text(lx,val_y,_compact(total),color=INK,fontsize=20,fontweight="bold",ha="left",va="center",zorder=4)
        bg.text(lx,sub_y,f"Latest {arrow} {abs(change):.0f}% vs avg",color=acol,fontsize=9,fontweight="bold",ha="left",va="center",zorder=4)
        bg.text(cxm,lab_y,"DAILY AVG",color=SUB,fontsize=8.5,ha="center",va="center",zorder=4)
        bg.text(cxm,val_y,_compact(avg),color=INK,fontsize=17,fontweight="bold",ha="center",va="center",zorder=4)
        bg.text(rx,lab_y,"PEAK",color=SUB,fontsize=8.5,ha="right",va="center",zorder=4)
        bg.text(rx,val_y,_compact(peak),color=INK,fontsize=17,fontweight="bold",ha="right",va="center",zorder=4)
        bg.text(rx,sub_y,f"({peakday})",color=SUB,fontsize=9,ha="right",va="center",zorder=4)

        # ---- bottom graph box ----
        bot=FancyBboxPatch((x0,BOT_Y0),colw,BOT_Y1-BOT_Y0,boxstyle="round,pad=0.002,rounding_size=0.012",
                           fc=CARD,ec=GRID,lw=1.2,zorder=2); bg.add_patch(bot)
        ax=fig.add_axes([x0+0.03, BOT_Y0+0.05, colw-0.048, (BOT_Y1-BOT_Y0)-0.072]); ax.set_facecolor(CARD)
        for xi,v in zip(x,vals):
            if v>0:
                ax.bar(xi,v,0.6,color=color,zorder=3)
                ax.annotate(_compact(v),(xi,v),textcoords="offset points",xytext=(0,3),ha="center",
                            color=INK,fontsize=8,fontweight="bold",clip_on=True)
        hr = 12 if i==0 else 6                                # extra headroom for Primary Sales
        ax.set_yscale("log"); ax.set_ylim(1, (max(pos)*hr if pos else 10))
        if not pos:
            ax.text(0.5,0.5,"No quantity data",transform=ax.transAxes,ha="center",va="center",
                    color=SUB,fontsize=11,style="italic")
        ax.set_xticks(x); ax.set_xticklabels([short(d) for d in dates],color=INK,fontsize=8)
        ax.tick_params(axis="y",colors=SUB,labelsize=7.5)
        for s in ax.spines.values(): s.set_color(GRID)
        ax.grid(axis="y",which="both",color="white",alpha=0.08,zorder=0)
    fig.savefig(str(out_png),dpi=dpi,facecolor=fig.get_facecolor()); plt.close(fig)
    print("Quantity PNG rendered:",out_png)
    return out_png

def _add_slide(prs, png):
    """Append the quantity slide (full-bleed PNG) to an existing Presentation."""
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(0x06,0x22,0x4A)
    slide.shapes.add_picture(str(png),0,0,width=prs.slide_width,height=prs.slide_height)
    return slide

def build(csv_paths, report_date, output, png=None):
    """Build a standalone quantity deck (one slide)."""
    data=_collect(csv_paths)
    tmp = Path(png) if png else (Path(output).with_suffix(".png"))
    render_png(data, report_date, tmp)
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    _add_slide(prs, tmp)
    prs.save(output)
    print("Saved quantity deck:",output)
    return output

def append_to_pptx(pptx_path, csv_paths, report_date, png=None):
    """Render the quantity chart and append it as a new slide to an EXISTING .pptx
    (the main dashboard deck), producing a single merged file in place."""
    data=_collect(csv_paths)
    tmp = Path(png) if png else (Path(pptx_path).with_name("_qty_tmp.png"))
    render_png(data, report_date, tmp)
    prs=Presentation(str(pptx_path))
    _add_slide(prs, tmp)
    prs.save(str(pptx_path))
    print("Appended quantity slide to:", pptx_path)
    return pptx_path

if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("--csv-file",nargs="+",required=True)
    ap.add_argument("--date",required=True)
    ap.add_argument("--output",default="",help="write a standalone one-slide quantity deck here")
    ap.add_argument("--append-to",default="",help="append the quantity slide to this existing .pptx")
    ap.add_argument("--png",default="")
    a=ap.parse_args()
    if a.append_to:
        append_to_pptx(a.append_to, a.csv_file, a.date, a.png or None)
    elif a.output:
        build(a.csv_file, a.date, a.output, a.png or None)
    else:
        ap.error("provide --output (standalone deck) or --append-to (merge into existing deck)")
